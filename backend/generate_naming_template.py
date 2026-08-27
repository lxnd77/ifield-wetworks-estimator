"""Generates a fill-in-the-blanks workbook for cleaning up line item / BOM
item naming and assigning purchase category, vendor, and purchasing company.

Two sheets, both keyed by database id so a filled-in copy can be re-imported
reliably even after names change:
  - Line Items: one row per WetworksProduct.
  - BOM Items: one row per BomLine (so a support item shared by several
    products appears once per product it's used in, with a "shared across N
    products" flag -- renaming it on any one row renames the single
    underlying record everywhere, per the agreed design). A block of blank
    rows is appended at the bottom for adding brand-new BOM items to a line
    item -- leave bom_line_id/support_item_id blank, pick the product from
    the dropdown, and fill in the rest; more rows can be added below those
    too, the importer doesn't care where a row sits.

Usage:  python generate_naming_template.py [output_path.xlsx]
"""
import sys
import os
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import joinedload

from app.database import SessionLocal
from app import models

PURCHASE_CATEGORIES = ["Paint", "Tile", "Stone", "Metal", "Other"]

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
EDITABLE_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
FLAG_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")


def style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def build(db, path):
    wb = openpyxl.Workbook()

    # ---------------------------------------------------------------- Instructions
    ws0 = wb.active
    ws0.title = "Instructions"
    lines = [
        "How to fill this out",
        "",
        "Two sheets: 'Line Items' (the Wetworks products themselves) and 'BOM Items' "
        "(the materials each one is built from).",
        "",
        "Yellow columns are editable. Everything else is read-only context -- don't need to touch it.",
        "",
        "Leave a 'New Name' cell blank to keep the current name as-is.",
        "",
        "On 'BOM Items': a red 'Name collision?' = YES means this BOM item currently has the exact same "
        "name as its own line item (the bug we're fixing) -- e.g. line item 'Floor GVT Tile 60 X120' whose "
        "BOM material is also literally named 'Floor GVT Tile 60 X120'. Give it a distinct material name, "
        "e.g. 'GVT Tile 60x120'.",
        "",
        "'Shared across N products' on 'BOM Items' means this exact material (screws, joint compound, etc.) "
        "is reused by more than one line item. It's a single database record -- renaming it on ANY row "
        "renames it everywhere it's used, so only rename it once and leave the other occurrences' 'New Name' "
        "blank.",
        "",
        "'Purchasing Company' (Line Items sheet) is the I-Field entity that buys the finished line item -- "
        "leave blank to keep whatever is already set (defaults to I-Field Dubai for Wetworks).",
        "",
        "'Direct Vendor Override' (Line Items sheet) is only for the rare case a line item is bought whole, "
        "directly from a vendor, bypassing the purchasing company. Leave blank unless that applies.",
        "",
        "'Vendor' (BOM Items sheet) is the actual external supplier that material is bought from.",
        "",
        "'Purchase Category' (BOM Items sheet) must be one of: " + ", ".join(PURCHASE_CATEGORIES) + " -- "
        "pick from the dropdown. This gates which items get a code field during estimation.",
        "",
        "Company/vendor names: type any name you like -- new ones will be created automatically on import. "
        "Already-known purchasing companies: I FIELD FURNISHING TRADING LLC (Dubai, default for Wetworks).",
        "",
        "Don't touch the 'id' columns -- they're how the re-import matches rows back to database records.",
        "",
        "ADDING A NEW BOM ITEM to an existing line item: go to the blank rows at the bottom of 'BOM Items' "
        "(marked '-- add new BOM items below --'), or just add a new row anywhere. Leave 'bom_line_id' and "
        "'support_item_id' blank -- that's what tells the importer this is a new BOM line, not an edit. Pick "
        "the line item it belongs to from the 'Product (context)' dropdown, put the material's name in 'New "
        "BOM Item Name', and fill in UoM, Qty per Unit (how much of it 1 unit of the product needs, before "
        "wastage), Wastage % (e.g. 10 for 10%, blank = 0), and Role (primary = the main material, fixing = "
        "everything else -- blank defaults to fixing). If the name you type already matches an existing "
        "material exactly, it reuses that same shared item instead of creating a duplicate -- e.g. attaching "
        "an existing screw or joint-compound item to another product this way is fine.",
    ]
    for i, text in enumerate(lines, start=1):
        cell = ws0.cell(row=i, column=1, value=text)
        if i == 1:
            cell.font = Font(bold=True, size=14)
    ws0.column_dimensions["A"].width = 120
    for i in range(1, len(lines) + 1):
        ws0.row_dimensions[i].height = 15 * (1 + len(lines[i - 1]) // 110) if lines[i - 1] else 6

    # ---------------------------------------------------------------- Line Items
    ws1 = wb.create_sheet("Line Items")
    headers1 = [
        "id", "Current Name", "New Name", "Category", "UoM",
        "Purchasing Company", "Direct Vendor Override",
    ]
    ws1.append(headers1)
    style_header(ws1, len(headers1))

    products = db.query(models.WetworksProduct).options(
        joinedload(models.WetworksProduct.purchasing_company),
        joinedload(models.WetworksProduct.default_vendor),
    ).order_by(models.WetworksProduct.category, models.WetworksProduct.name).all()

    for p in products:
        ws1.append([
            p.id, p.name, "", p.category, p.uom,
            p.purchasing_company.name if p.purchasing_company else "",
            p.default_vendor.name if p.default_vendor else "",
        ])
    for row in range(2, ws1.max_row + 1):
        ws1.cell(row=row, column=3).fill = EDITABLE_FILL
        ws1.cell(row=row, column=6).fill = EDITABLE_FILL
        ws1.cell(row=row, column=7).fill = EDITABLE_FILL
    autosize(ws1, [6, 42, 42, 16, 8, 30, 30])

    # ---------------------------------------------------------------- BOM Items
    ws2 = wb.create_sheet("BOM Items")
    headers2 = [
        "bom_line_id", "support_item_id", "Product (context)", "Current BOM Item Name", "New BOM Item Name",
        "Name collision?", "Shared across N products", "UoM", "Purchase Category", "Vendor",
        "Qty per Unit", "Wastage %", "Role",
    ]
    ws2.append(headers2)
    style_header(ws2, len(headers2))

    bom_lines = db.query(models.BomLine).options(
        joinedload(models.BomLine.product),
        joinedload(models.BomLine.support_item).joinedload(models.SupportItem.default_vendor),
    ).join(models.WetworksProduct).order_by(
        models.WetworksProduct.category, models.WetworksProduct.name, models.BomLine.sort_order
    ).all()

    usage_count = {}
    for b in bom_lines:
        usage_count[b.support_item_id] = usage_count.get(b.support_item_id, 0) + 1

    for b in bom_lines:
        si = b.support_item
        collision = si.name.strip().lower() == b.product.name.strip().lower()
        shared_n = usage_count[si.id]
        ws2.append([
            b.id, si.id, b.product.name, si.name, "",
            "YES" if collision else "no",
            shared_n if shared_n > 1 else "",
            si.uom, si.purchase_category or "",
            si.default_vendor.name if si.default_vendor else "",
            b.qty_per_unit, round((b.wastage_pct or 0) * 100, 2), b.role,
        ])
        r = ws2.max_row
        if collision:
            ws2.cell(row=r, column=6).fill = FLAG_FILL

    existing_row_count = ws2.max_row

    # Separator + blank rows for adding brand-new BOM items.
    ws2.append(["", "", "-- add new BOM items below --", "", "", "", "", "", "", "", "", "", ""])
    sep_row = ws2.max_row
    for col in range(1, len(headers2) + 1):
        ws2.cell(row=sep_row, column=col).fill = PatternFill(start_color="D1D5DB", end_color="D1D5DB", fill_type="solid")
        ws2.cell(row=sep_row, column=col).font = Font(italic=True)

    NEW_ROW_BLOCK = 20
    for _ in range(NEW_ROW_BLOCK):
        ws2.append(["", "", "", "", "", "", "", "", "", "", "", "", ""])

    last_row = ws2.max_row

    for row in list(range(2, existing_row_count + 1)) + list(range(sep_row + 1, last_row + 1)):
        for col in (3, 5, 8, 9, 10, 11, 12, 13):  # Product, New Name, UoM, Category, Vendor, Qty, Wastage, Role
            ws2.cell(row=row, column=col).fill = EDITABLE_FILL

    category_dv = DataValidation(type="list", formula1='"' + ",".join(PURCHASE_CATEGORIES) + '"', allow_blank=True)
    ws2.add_data_validation(category_dv)
    category_dv.add(f"I2:I{last_row}")

    role_dv = DataValidation(type="list", formula1='"primary,fixing"', allow_blank=True)
    ws2.add_data_validation(role_dv)
    role_dv.add(f"M2:M{last_row}")

    product_dv = DataValidation(type="list", formula1=f"='Line Items'!$B$2:$B${ws1.max_row}", allow_blank=True)
    ws2.add_data_validation(product_dv)
    product_dv.add(f"C2:C{last_row}")

    autosize(ws2, [10, 14, 42, 42, 42, 14, 20, 8, 16, 30, 12, 10, 10])

    wb.save(path)
    print(f"Wrote {path}")
    print(f"  Line Items: {len(products)} rows")
    print(f"  BOM Items: {len(bom_lines)} rows ({sum(1 for b in bom_lines if b.support_item.name.strip().lower() == b.product.name.strip().lower())} name collisions flagged)")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "naming_template.xlsx"
    db = SessionLocal()
    try:
        build(db, out)
    finally:
        db.close()
