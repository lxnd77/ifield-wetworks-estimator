"""Supplier BOQ: a presentable, I-Field-branded bill of quantities to send to
suppliers for pricing. Unlike the three Odoo exports it carries no costs,
margins or vendors, and it works on an *unfinished* estimate -- nothing is
validated, blank item codes / quantities just come out blank.

Sheets:
  - "BOQ": one row per estimate line (location, item code, item,
    description, dimension, drawing no, unit, qty, remark).
  - One sheet per BOM purchase category (furniture: Fabric / Stone / Metal /
    Accessories, then Factory Work; wetworks: whatever categories its recipe
    items carry), each item totalled across the project by (support item,
    item code) so a category can go to its own supplier.

Every sheet ends in blank shaded Unit Price cells, a Total formula and a
grand total, for the supplier to fill in and return.
"""
import re
from datetime import date
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import models, project_types, uoms
from .export_excel import norm_code, reference_code

# Brand palette (I-Field guidelines). Georgia / Arial stand in for Playfair
# Display / Montserrat, which a supplier's Excel usually doesn't have.
BLACK = "231F20"
RUBY = "991B1E"
CREAM = "F8F5F0"
RULE = "D9D4CC"
HEAD_FONT = "Georgia"
BODY_FONT = "Arial"

LOGO_PATH = Path(__file__).resolve().parent / "assets" / "ifield_logo.png"

OTHER_CATEGORY = "Other"
FACTORY_WORK_SHEET = "Factory Work"

_thin = Side(style="thin", color=RULE)
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_QTY_FMT = "#,##0.00"
_PRICE_FMT = "#,##0.00"

# (header, width, kind) -- kind drives formatting: "qty", "price", "total"
# or None for text.
LINE_COLUMNS = [
    ("No.", 6, None), ("Location", 18, None), ("Item Code", 14, None), ("Item", 34, None),
    ("Description", 30, None), ("Dimension", 18, None), ("Drawing No.", 14, None),
    ("Unit", 9, None), ("Qty", 10, "qty"), ("Unit Price", 13, "price"), ("Total", 15, "total"),
    ("Remarks", 28, None),
]
ITEM_COLUMNS = [
    ("No.", 6, None), ("Item Code", 14, None), ("Item", 38, None), ("Unit", 9, None),
    ("Total Qty", 11, "qty"), ("Used In", 32, None), ("Unit Price", 13, "price"),
    ("Total", 15, "total"), ("Remarks", 28, None),
]


def _sheet_title(name: str, used: set) -> str:
    """Excel sheet names: <= 31 chars, none of []:*?/\\, unique."""
    base = re.sub(r"[\[\]:*?/\\]", "-", name).strip()[:31] or "Sheet"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _add_logo(ws) -> bool:
    try:
        from openpyxl.drawing.image import Image
        img = Image(str(LOGO_PATH))
    except Exception:  # no Pillow / no file: fall back to the text title only
        return False
    img.height, img.width = 36, round(36 * img.width / img.height)
    ws.add_image(img, "A1")
    return True


def _write_header(ws, project: models.Project, subtitle: str, ncols: int) -> int:
    """The title block shared by every sheet. Returns the table's header row."""
    last = get_column_letter(ncols)
    row = 3 if _add_logo(ws) else 1
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 12

    ws.cell(row=row, column=1, value="Bill of Quantities").font = Font(name=HEAD_FONT, size=18, bold=True, color=BLACK)
    ws.row_dimensions[row].height = 26
    row += 1
    ws.cell(row=row, column=1, value=subtitle.upper()).font = Font(name=BODY_FONT, size=9, bold=True, color=RUBY)
    row += 2

    project_label = project.name + (f" ({project.code})" if project.code else "")
    info = [
        ("Project", project_label, "Supplier", None),
        ("Client", project.client_name or "", "Quoted currency", None),
        ("Site", project.address or (project.country.name if project.country else ""), "Quote valid until", None),
        ("Issued", date.today().strftime("%d %b %Y"), "Contact", None),
    ]
    label_font = Font(name=BODY_FONT, size=9, color="7A746C")
    value_font = Font(name=BODY_FONT, size=10, color=BLACK)
    # Project facts on the left (label in A, value from B); the supplier's
    # fields on the right: a label over two columns, then a shaded input
    # over the last three.
    right_label_col = ncols - 4
    input_col = ncols - 2
    for left_label, left_value, right_label, _ in info:
        ws.cell(row=row, column=1, value=left_label).font = label_font
        ws.cell(row=row, column=2, value=left_value).font = value_font
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=right_label_col - 1)
        label = ws.cell(row=row, column=right_label_col, value=right_label)
        label.font = label_font
        label.alignment = Alignment(horizontal="right", indent=1)
        ws.merge_cells(start_row=row, start_column=right_label_col, end_row=row, end_column=input_col - 1)
        for col in range(input_col, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = PatternFill("solid", fgColor=CREAM)
            cell.border = _BORDER
        ws.merge_cells(start_row=row, start_column=input_col, end_row=row, end_column=ncols)
        row += 1

    row += 1
    note = ws.cell(row=row, column=1,
                   value="Please enter your unit prices in the shaded cells, fill in your details above, "
                         "and return this file. Quantities are for the whole project.")
    note.font = Font(name=BODY_FONT, size=9, italic=True, color="5A554F")
    ws.merge_cells(f"A{row}:{last}{row}")
    return row + 2


def _write_table(ws, project, subtitle, columns, rows) -> None:
    """rows: lists of values for every column except No., Unit Price and Total
    (filled here)."""
    ncols = len(columns)
    header_row = _write_header(ws, project, subtitle, ncols)
    kinds = [k for _, _, k in columns]
    qty_col = kinds.index("qty") + 1
    price_col = kinds.index("price") + 1
    total_col = kinds.index("total") + 1

    for i, (title, width, _) in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=i, value=title)
        c.font = Font(name=BODY_FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RUBY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 22

    body_font = Font(name=BODY_FONT, size=9, color=BLACK)
    r = header_row
    for n, values in enumerate(rows, start=1):
        r += 1
        values = list(values)
        full = [n]
        for kind in kinds[1:]:
            full.append(None if kind in ("price", "total") else values.pop(0))
        for col, value in enumerate(full, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = body_font
            c.border = _BORDER
            kind = kinds[col - 1]
            if kind == "qty":
                c.number_format = _QTY_FMT
                c.alignment = Alignment(horizontal="right", vertical="top")
            elif kind in ("price", "total"):
                c.number_format = _PRICE_FMT
                c.alignment = Alignment(horizontal="right", vertical="top")
            elif col == 1:
                c.alignment = Alignment(horizontal="center", vertical="top")
            else:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=price_col).fill = PatternFill("solid", fgColor=CREAM)
        q = f"{get_column_letter(qty_col)}{r}"
        p = f"{get_column_letter(price_col)}{r}"
        ws.cell(row=r, column=total_col, value=f'=IF({p}="","",{q}*{p})')

    if not rows:
        r += 1
        c = ws.cell(row=r, column=2, value="No items yet.")
        c.font = Font(name=BODY_FONT, size=9, italic=True, color="7A746C")

    r += 1
    label = ws.cell(row=r, column=price_col, value="Total")
    label.font = Font(name=BODY_FONT, size=10, bold=True, color=BLACK)
    label.alignment = Alignment(horizontal="right")
    tl = get_column_letter(total_col)
    total = ws.cell(row=r, column=total_col,
                    value=f"=SUM({tl}{header_row + 1}:{tl}{r - 1})" if r - 1 > header_row else 0)
    total.font = Font(name=BODY_FONT, size=10, bold=True, color=RUBY)
    total.number_format = _PRICE_FMT
    top = Border(top=Side(style="medium", color=RUBY))
    for col in (price_col, total_col):
        ws.cell(row=r, column=col).border = top

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.center.size = 8
    ws.sheet_view.showGridLines = False


def _ordered_lines(project: models.Project) -> list:
    """Lines in estimator order: by location (sort order, then creation), then line."""
    def key(line):
        loc = line.location
        return ((loc.sort_order or 0) if loc else 0, loc.id if loc else 0, line.id or 0)
    return sorted(project.estimate_lines, key=key)


def _joined(values) -> str:
    seen = []
    for v in values:
        if v and v not in seen:
            seen.append(v)
    return ", ".join(seen)


def build_supplier_boq_workbook(project: models.Project) -> BytesIO:
    furniture = project_types.bom_per_line(project.project_type)
    lines = _ordered_lines(project)
    wb = openpyxl.Workbook()
    used_titles = set()

    ws = wb.active
    ws.title = _sheet_title("BOQ", used_titles)
    _write_table(ws, project, f"{project_types.label(project.project_type)} -- line items", LINE_COLUMNS, [
        [
            line.location.name if line.location else "",
            reference_code(project, line.item_code),
            line.product.name,
            line.description or "",
            line.dimension or "",
            line.drawing_no or "",
            uoms.normalize(line.product.uom),
            line.qty,
            line.remark or "",
        ]
        for line in lines
    ])

    # Support items totalled per (item, code), bucketed by purchase category.
    by_category: dict = {}
    for line in lines:
        line_ref = reference_code(project, line.item_code)
        for comp in line.components:
            si = comp.support_item
            category = (si.purchase_category or "").strip() or OTHER_CATEGORY
            bucket = by_category.setdefault(category, {})
            key = (comp.support_item_id, norm_code(comp.item_code))
            entry = bucket.setdefault(key, {
                "code": reference_code(project, comp.item_code), "name": si.name,
                "uom": uoms.normalize(si.uom), "qty": 0.0, "used_in": [],
            })
            entry["qty"] += comp.qty or 0.0
            entry["used_in"].append(line_ref or line.product.name)

    if furniture:
        order = list(project_types.FURNITURE_BOM_CATEGORIES)
    else:
        order = []
    order += sorted(c for c in by_category if c not in order and c != OTHER_CATEGORY)
    order.append(OTHER_CATEGORY)

    for category in order:
        bucket = by_category.get(category)
        if not bucket:
            continue
        # Wetworks recipe items mostly carry no category: "Materials" reads
        # better than "Other" when that's the only BOM sheet.
        if category == OTHER_CATEGORY and len(by_category) == 1:
            category = "Materials"
        ws = wb.create_sheet(_sheet_title(category, used_titles))
        items = sorted(bucket.values(), key=lambda e: (e["name"].lower(), e["code"].lower()))
        _write_table(ws, project, f"{category} -- totalled across the project", ITEM_COLUMNS, [
            [e["code"], e["name"], e["uom"], round(e["qty"], 4), _joined(e["used_in"]), ""]
            for e in items
        ])

    if furniture and lines:
        # One Factory Work row per exported product (product + item code),
        # qty = units across every line of it.
        fw: dict = {}
        for line in lines:
            key = (line.product_id, norm_code(line.item_code))
            entry = fw.setdefault(key, {
                "code": reference_code(project, line.item_code),
                "name": f"Factory Work for {line.product.name}", "qty": 0.0, "used_in": [],
            })
            entry["qty"] += line.qty or 0.0
            entry["used_in"].append(line.location.name if line.location else "")
        ws = wb.create_sheet(_sheet_title(FACTORY_WORK_SHEET, used_titles))
        _write_table(ws, project, "Factory Work -- per finished product", ITEM_COLUMNS, [
            [e["code"], e["name"], uoms.FACTORY_WORK, round(e["qty"], 4), _joined(e["used_in"]), ""]
            for e in fw.values()
        ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
