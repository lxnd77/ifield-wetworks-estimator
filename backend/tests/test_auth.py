"""Auth + project-isolation tests. Run with: pytest tests/ (from backend/)."""
import os
import sys
import uuid
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

# Point at an isolated on-disk SQLite file *before* importing anything from
# `app` -- database.py reads DATABASE_URL at import time.
_TEST_DB = os.path.join(os.path.dirname(__file__), f"test_auth_{uuid.uuid4().hex}.db")
os.environ["DATABASE_URL"] = f"sqlite:///{_TEST_DB}"

import pytest
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal, engine
from app import models, service, export_excel
from app.auth import hash_password


@pytest.fixture(scope="module")
def client():
    db = SessionLocal()
    country = models.Country(name="Testland", code="TST", is_active=True, is_template=False)
    db.add(country)
    db.flush()
    db.add(models.User(username="admin", password_hash=hash_password("adminpass"), is_admin=True))
    db.add(models.User(username="alice", password_hash=hash_password("alicepass"), is_admin=False))
    db.add(models.User(username="bob", password_hash=hash_password("bobpass"), is_admin=False))
    db.commit()
    db.close()

    with TestClient(app) as c:
        yield c

    engine.dispose()
    if os.path.exists(_TEST_DB):
        os.remove(_TEST_DB)


def _login(client, username, password):
    r = client.post("/api/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def test_login_success(client):
    token = _login(client, "alice", "alicepass")
    r = client.get("/api/auth/me", headers=_auth(token))
    assert r.status_code == 200
    assert r.json()["username"] == "alice"


def test_login_wrong_password(client):
    r = client.post("/api/auth/login", json={"username": "alice", "password": "wrong"})
    assert r.status_code == 401


def test_login_unknown_user(client):
    r = client.post("/api/auth/login", json={"username": "nobody", "password": "x"})
    assert r.status_code == 401


def test_missing_token_rejected(client):
    r = client.get("/api/countries")
    assert r.status_code == 401


def test_garbage_token_rejected(client):
    r = client.get("/api/countries", headers=_auth("not-a-real-token"))
    assert r.status_code == 401


def test_valid_token_allows_access(client):
    token = _login(client, "alice", "alicepass")
    r = client.get("/api/countries", headers=_auth(token))
    assert r.status_code == 200


def _new_project(client, token, name):
    r = client.post("/api/projects", headers=_auth(token), json={
        "name": name, "code": "NP", "country_id": _country_id(client, token),
        "start_date": "2026-01-01", "end_date": "2026-06-01",
    })
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_project_isolation_between_users(client):
    alice_token = _login(client, "alice", "alicepass")
    bob_token = _login(client, "bob", "bobpass")

    alice_project_id = _new_project(client, alice_token, "Alice's Project")
    bob_project_id = _new_project(client, bob_token, "Bob's Project")

    alice_list = client.get("/api/projects", headers=_auth(alice_token)).json()
    assert [p["name"] for p in alice_list] == ["Alice's Project"]

    bob_list = client.get("/api/projects", headers=_auth(bob_token)).json()
    assert [p["name"] for p in bob_list] == ["Bob's Project"]

    # Bob can't see or touch Alice's project by id.
    assert client.get(f"/api/projects/{alice_project_id}", headers=_auth(bob_token)).status_code == 404
    assert client.delete(f"/api/projects/{alice_project_id}", headers=_auth(bob_token)).status_code == 404
    assert client.post(f"/api/projects/{alice_project_id}/locations", headers=_auth(bob_token),
                        json={"name": "Lobby"}).status_code == 404


def test_admin_sees_all_projects(client):
    admin_token = _login(client, "admin", "adminpass")
    names = {p["name"] for p in client.get("/api/projects", headers=_auth(admin_token)).json()}
    assert {"Alice's Project", "Bob's Project"}.issubset(names)


def test_non_admin_cannot_manage_users(client):
    alice_token = _login(client, "alice", "alicepass")
    assert client.get("/api/users", headers=_auth(alice_token)).status_code == 403
    assert client.post("/api/users", headers=_auth(alice_token),
                        json={"username": "eve", "password": "x"}).status_code == 403


def test_admin_can_create_user(client):
    admin_token = _login(client, "admin", "adminpass")
    r = client.post("/api/users", headers=_auth(admin_token), json={"username": "carol", "password": "carolpass"})
    assert r.status_code == 200, r.text
    assert r.json()["is_admin"] is False
    # New account can log in immediately.
    _login(client, "carol", "carolpass")


def test_admin_can_reset_password(client):
    admin_token = _login(client, "admin", "adminpass")
    dave_id = client.post("/api/users", headers=_auth(admin_token),
                           json={"username": "dave", "password": "oldpass"}).json()["id"]
    r = client.put(f"/api/users/{dave_id}/password", headers=_auth(admin_token), json={"new_password": "newpass"})
    assert r.status_code == 200, r.text
    _login(client, "dave", "newpass")
    assert client.post("/api/auth/login", json={"username": "dave", "password": "oldpass"}).status_code == 401


def test_non_admin_cannot_reset_password(client):
    alice_token = _login(client, "alice", "alicepass")
    admin_id = client.get("/api/auth/me", headers=_auth(alice_token)).json()["id"]
    r = client.put(f"/api/users/{admin_id}/password", headers=_auth(alice_token), json={"new_password": "x"})
    assert r.status_code == 403


# ---- project types (furniture extension, phase 01) ----

def test_project_types_endpoint(client):
    token = _login(client, "alice", "alicepass")
    r = client.get("/api/project-types", headers=_auth(token))
    assert r.status_code == 200, r.text
    by_value = {t["value"]: t for t in r.json()}
    assert set(by_value) == {"wetworks", "loose_furniture", "fixed_furniture"}
    assert by_value["wetworks"]["labor_applies"] is True
    assert by_value["loose_furniture"]["labor_applies"] is False
    assert by_value["fixed_furniture"]["dates_required"] is False


def _country_id(client, token):
    countries = client.get("/api/countries", headers=_auth(token)).json()
    return next(c["id"] for c in countries if c["code"] == "TST")


def test_project_code_required(client):
    token = _login(client, "alice", "alicepass")
    for code in (None, "", "   "):
        r = client.post("/api/projects", headers=_auth(token), json={
            "name": "No code", "code": code, "country_id": _country_id(client, token),
            "project_type": "loose_furniture",
        })
        assert r.status_code == 422 and "project code" in r.json()["detail"]


def test_furniture_project_allows_missing_dates(client):
    token = _login(client, "alice", "alicepass")
    r = client.post("/api/projects", headers=_auth(token), json={
        "name": "LF Project", "code": "LF", "country_id": _country_id(client, token),
        "project_type": "loose_furniture",
    })
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["project_type"] == "loose_furniture"
    assert body["start_date"] is None and body["end_date"] is None


def test_wetworks_project_requires_dates(client):
    token = _login(client, "alice", "alicepass")
    r = client.post("/api/projects", headers=_auth(token), json={
        "name": "WW no dates", "code": "WW", "country_id": _country_id(client, token),
    })
    assert r.status_code == 422


def test_unknown_project_type_rejected(client):
    token = _login(client, "alice", "alicepass")
    r = client.post("/api/projects", headers=_auth(token), json={
        "name": "bogus", "country_id": _country_id(client, token),
        "project_type": "nope", "start_date": "2026-01-01", "end_date": "2026-02-01",
    })
    assert r.status_code == 422


def test_products_filtered_by_type(client):
    token = _login(client, "alice", "alicepass")
    # test DB has no products at all, but the filter must still be honored
    # (and not error) -- an unknown type simply yields an empty list.
    r = client.get("/api/products?product_type=loose_furniture", headers=_auth(token))
    assert r.status_code == 200
    assert r.json() == []


# ---- furniture costing (phase 02 / 05 / 08) ----

_line_seq = [0]


def _make_line(db, project_type, *, with_coverage=False, factory_work_cny=None, item_code=None,
               comp_code=None, comp_qty_per_unit=2.0, comp_price_cny=10.0,
               ohp_pct=0.0, cny_per_usd=1.0, project_code=None):
    """Insert a self-contained catalog + a one-line estimate; return the line.
    Testland fx (for wetworks) = 1.0, line qty = 3.

    Wetworks: a fixed BomLine recipe (2 per unit of a $10 support item).
    Furniture: a user-authored EstimateLineComponent (support item +
    comp_qty_per_unit + comp_price_cny (CNY) + comp_code), plus the per-line
    factory_work_cny charge. cny_per_usd defaults to 1.0 so CNY == USD.
    """
    country = db.query(models.Country).filter_by(code="TST").one()
    _line_seq[0] += 1
    suffix = f"{project_type}-{_line_seq[0]}"
    furniture = project_type != "wetworks"
    product = models.Product(
        name=f"P {suffix}", uom="Pcs", category="Test", product_type=project_type,
        needs_setup=False, consumable_pct=0.0, ohp_pct=ohp_pct,
    )
    db.add(product)
    db.flush()
    si = models.SupportItem(name=f"SI {suffix}", uom="Pcs", purchase_category="Fabric")
    db.add(si)
    db.flush()
    if not furniture:
        db.add(models.CountryMaterialPrice(country_id=country.id, support_item_id=si.id,
                                           unit_price_local=10.0))
        db.add(models.BomLine(product_id=product.id, support_item_id=si.id,
                              qty_per_unit=2.0, wastage_pct=0.0, role="primary", sort_order=0))
    if with_coverage:
        db.add(models.CoverageRate(
            product_id=product.id, primary_coverage_per_day=10.0, secondary_coverage_per_day=0.0,
            inhouse_count=1, local_count=0, inhouse_salary_month_local=2600.0,
            local_salary_month_local=0.0))
    project = models.Project(
        name=f"Proj {suffix}", country_id=country.id, project_type=project_type, code=project_code,
        start_date=date(2026, 1, 1), end_date=date(2026, 7, 1), default_margin_pct=0.0,
        cny_per_usd=cny_per_usd if furniture else None)
    db.add(project)
    db.flush()
    loc = models.ProjectLocation(project_id=project.id, name="L")
    db.add(loc)
    db.flush()
    line = models.EstimateLine(
        project_id=project.id, location_id=loc.id, product_id=product.id, qty=3.0,
        item_code=item_code, factory_work_cost_cny=factory_work_cny)
    db.add(line)
    db.flush()
    if furniture:
        db.add(models.EstimateLineComponent(
            estimate_line_id=line.id, support_item_id=si.id, qty_per_unit=comp_qty_per_unit,
            unit_price_cny=comp_price_cny, item_code=comp_code or f"C{_line_seq[0]}",
            qty=0.0, unit_cost=0.0, total_cost=0.0))
        db.flush()
    return line


def test_furniture_line_prices_material_only():
    db = SessionLocal()
    try:
        # Even with a stray coverage rate, a furniture line carries no labor;
        # material = components (2*10) + Factory Work (5), fx 1.0, ohp 0.
        line = _make_line(db, "loose_furniture", with_coverage=True,
                          factory_work_cny=5.0, item_code="F1")
        service.recompute_estimate_line(db, line)
        db.commit()
        assert line.material_cost_per_unit == 25.0
        assert line.labor_cost_per_unit == 0.0
        comp = line.components[0]
        assert comp.qty == 6.0 and comp.unit_cost == 10.0 and comp.total_cost == 60.0
    finally:
        db.close()


def test_furniture_ohp_applies_to_line_total_not_components():
    db = SessionLocal()
    try:
        line = _make_line(db, "fixed_furniture", factory_work_cny=10.0, item_code="F2", ohp_pct=0.10)
        service.recompute_estimate_line(db, line)
        db.commit()
        # (2*10 + 10) * 1.10 = 33
        assert abs(line.material_cost_per_unit - 33.0) < 1e-9
        # component total stays raw (no OHP): 2 * 10 * qty(3)
        assert line.components[0].total_cost == 60.0
    finally:
        db.close()


def test_furniture_cny_conversion():
    db = SessionLocal()
    try:
        line = _make_line(db, "loose_furniture", cny_per_usd=7.2, comp_qty_per_unit=1.0,
                          comp_price_cny=72.0, factory_work_cny=36.0, item_code="F-CNY")
        service.recompute_estimate_line(db, line)
        db.commit()
        assert abs(line.components[0].unit_cost - 10.0) < 1e-9      # 72 / 7.2
        assert abs(line.material_cost_per_unit - 15.0) < 1e-9        # 1*10 + 36/7.2
    finally:
        db.close()


def test_furniture_component_reprices_on_qty_change():
    db = SessionLocal()
    try:
        line = _make_line(db, "loose_furniture", factory_work_cny=1.0, item_code="F3")
        service.recompute_estimate_line(db, line)
        db.commit()
        assert line.material_cost_per_unit == 21.0
        line.qty = 5.0
        service.recompute_estimate_line(db, line)
        db.commit()
        assert line.components[0].qty == 10.0          # 2 per unit * 5
        assert line.material_cost_per_unit == 21.0     # per-unit cost unchanged
    finally:
        db.close()


def test_wetworks_line_still_has_labor():
    db = SessionLocal()
    try:
        line = _make_line(db, "wetworks", with_coverage=True)
        service.recompute_estimate_line(db, line)
        db.commit()
        assert line.material_cost_per_unit == 20.0
        assert abs(line.labor_cost_per_unit - 11.5) < 1e-6
    finally:
        db.close()


# ---- exports: estimation_type_id + labor column per type (phase 03) ----

def _sheet_rows(buf):
    ws = openpyxl.load_workbook(buf).active
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


def _built_line(db, project_type, *, with_coverage=False, item_code="X1", project_code="PJ", ohp_pct=0.0):
    line = _make_line(db, project_type, with_coverage=with_coverage, item_code=item_code,
                      project_code=project_code, ohp_pct=ohp_pct,
                      factory_work_cny=7.0 if project_type != "wetworks" else None)
    service.recompute_estimate_line(db, line)
    db.commit()
    return line


def test_sale_estimation_furniture_qualified_names_ohp_and_blank_labor():
    db = SessionLocal()
    try:
        line = _built_line(db, "loose_furniture", with_coverage=True, ohp_pct=0.08)
        header, data = _sheet_rows(export_excel.build_sale_estimation_workbook(db, line.project))
        assert data[0][header.index("estimation_type_id")] == "Loose Furniture"
        assert all(r[header.index("estimation_line_ids/labor_cost")] in (None, "") for r in data)
        # OHP handed to Odoo as a line-level percentage
        ohp = [r[header.index("estimation_line_ids/overhead_cost_percentage")] for r in data
               if r[header.index("estimation_line_ids/overhead_cost_percentage")] not in (None, "")]
        assert ohp == [0.08]
        # names qualified with project code + item code
        comp = line.components[0]
        pid_col = header.index("estimation_line_ids/product_id")
        comp_col = header.index(
            "estimation_line_ids/sale_estimation_component_product_line_ids/product_id")
        assert data[0][pid_col] == f"{line.product.name} PJ {line.item_code}"
        names = [r[comp_col] for r in data if r[comp_col]]
        assert f"{comp.support_item.name} PJ {comp.item_code}" in names
        assert f"Factory Work for {line.product.name} PJ {line.item_code}" in names
    finally:
        db.close()


def test_sale_estimation_wetworks_type_and_labor_present():
    db = SessionLocal()
    try:
        line = _built_line(db, "wetworks", with_coverage=True)
        header, data = _sheet_rows(export_excel.build_sale_estimation_workbook(db, line.project))
        assert data[0][header.index("estimation_type_id")] == "Wetworks"
        labor_col = header.index("estimation_line_ids/labor_cost")
        labor_vals = [r[labor_col] for r in data if r[labor_col] not in (None, "")]
        assert labor_vals and abs(labor_vals[0] - 11.5) < 1e-6
    finally:
        db.close()


def test_bom_export_furniture_qualified_names_and_factory_work():
    db = SessionLocal()
    try:
        line = _built_line(db, "fixed_furniture")
        bom_header, bom_data = _sheet_rows(export_excel.build_bom_workbook(db, line.project))
        comp = line.components[0]
        assert bom_data[0][bom_header.index("product")] == f"{line.product.name} PJ {line.item_code}"
        fw_name = f"Factory Work for {line.product.name} PJ {line.item_code}"
        names = [r[bom_header.index("bom_line_ids/product_id")] for r in bom_data]
        assert f"{comp.support_item.name} PJ {comp.item_code}" in names
        assert fw_name in names
        fw_qty = [r[bom_header.index("bom_line_ids/product_qty")] for r in bom_data
                  if r[bom_header.index("bom_line_ids/product_id")] == fw_name]
        assert fw_qty == [1]
    finally:
        db.close()


def test_product_import_furniture_manufacture_and_factory_work_buy_row():
    db = SessionLocal()
    try:
        line = _built_line(db, "loose_furniture")
        sc = models.SellingCompany(name="SC ftest")
        pc = models.PurchasingCompany(name="PC ftest")
        db.add_all([sc, pc])
        db.flush()
        line.project.selling_company_id = sc.id
        line.product.purchasing_company_id = pc.id
        db.commit()
        books = dict(export_excel.build_product_import_workbooks(db, line.project))
        _, sc_rows = _sheet_rows(books["SC ftest"])
        by_name = {r[1]: r for r in sc_rows}
        assert by_name[f"{line.product.name} PJ {line.item_code}"][4].startswith("Manufacture")
        comp = line.components[0]
        comp_row = by_name[f"{comp.support_item.name} PJ {comp.item_code}"]
        assert comp_row[4].startswith("Buy") and comp_row[8] == round(comp.unit_cost, 4)
        fw = by_name[f"Factory Work for {line.product.name} PJ {line.item_code}"]
        assert fw[4].startswith("Buy") and fw[7] == "Storable Product" and fw[8] in (None, "")
    finally:
        db.close()


def _furniture_project(client, h, cid):
    prod = client.post("/api/products", headers=h, json={
        "name": "F chair", "uom": "Pcs", "category": "Seating",
        "product_type": "loose_furniture"}).json()
    si = client.post("/api/support-items", headers=h, json={
        "name": "F fabric", "uom": "m", "purchase_category": "Fabric"}).json()["id"]
    pid = client.post("/api/projects", headers=h, json={
        "name": "F proj", "country_id": cid, "project_type": "loose_furniture", "code": "FP"}).json()["id"]
    loc = client.post(f"/api/projects/{pid}/locations", headers=h, json={"name": "L"}).json()["id"]
    return prod["id"], si, pid, loc


def test_furniture_component_requires_code_and_price(client):
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    prod, si, pid, loc = _furniture_project(client, h, _country_id(client, token))
    ln = client.post(f"/api/projects/{pid}/estimate-lines", headers=h,
                     json={"location_id": loc, "product_id": prod, "qty": 2}).json()["id"]
    # missing item_code -> 422 (schema)
    assert client.post(f"/api/estimate-lines/{ln}/components", headers=h,
                       json={"support_item_id": si, "qty_per_unit": 3, "unit_price_cny": 50}).status_code == 422
    # price <= 0 -> 422
    assert client.post(f"/api/estimate-lines/{ln}/components", headers=h,
                       json={"support_item_id": si, "qty_per_unit": 3, "unit_price_cny": 0,
                             "item_code": "FAB-1"}).status_code == 422
    # good -> 200
    r = client.post(f"/api/estimate-lines/{ln}/components", headers=h,
                    json={"support_item_id": si, "qty_per_unit": 3, "unit_price_cny": 50, "item_code": "FAB-1"})
    assert r.status_code == 200, r.text
    assert r.json()["unit_price_cny"] == 50


def test_furniture_component_code_is_item_plus_code(client):
    """Uniqueness is support item + code: the same code on a different item is
    fine; the same item + code elsewhere must keep the same price."""
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    prod, si, pid, loc = _furniture_project(client, h, _country_id(client, token))
    si2 = client.post("/api/support-items", headers=h, json={
        "name": "F stone", "uom": "Sqm", "purchase_category": "Stone"}).json()["id"]
    l1 = client.post(f"/api/projects/{pid}/estimate-lines", headers=h,
                     json={"location_id": loc, "product_id": prod, "qty": 1}).json()["id"]
    assert client.post(f"/api/estimate-lines/{l1}/components", headers=h,
                       json={"support_item_id": si, "qty_per_unit": 1, "unit_price_cny": 10,
                             "item_code": "DUP"}).status_code == 200
    l2 = client.post(f"/api/projects/{pid}/estimate-lines", headers=h,
                     json={"location_id": loc, "product_id": prod, "qty": 1}).json()["id"]
    # different support item, same code -> allowed
    assert client.post(f"/api/estimate-lines/{l2}/components", headers=h,
                       json={"support_item_id": si2, "qty_per_unit": 1, "unit_price_cny": 10,
                             "item_code": "dup"}).status_code == 200
    # same support item + code at the same price -> allowed (same product)
    assert client.post(f"/api/estimate-lines/{l2}/components", headers=h,
                       json={"support_item_id": si, "qty_per_unit": 2, "unit_price_cny": 10,
                             "item_code": "dup"}).status_code == 200
    comp_id = client.get(f"/api/projects/{pid}/estimate-lines", headers=h).json()[1]["components"][1]["id"]
    # ...but not at a different price
    r = client.put(f"/api/estimate-line-components/{comp_id}", headers=h,
                   json={"support_item_id": si, "qty_per_unit": 2, "unit_price_cny": 12, "item_code": "dup"})
    assert r.status_code == 422 and "already priced" in r.json()["detail"]


def test_new_furniture_project_snapshots_cny_rate(client):
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    cid = _country_id(client, token)
    r = client.post("/api/projects", headers=h, json={
        "name": "fx proj", "code": "FX", "country_id": cid, "project_type": "loose_furniture"})
    assert r.status_code == 200
    assert r.json()["cny_per_usd"] == 7.2
    ww = client.post("/api/projects", headers=h, json={
        "name": "ww fx", "code": "WF", "country_id": cid, "start_date": "2026-01-01", "end_date": "2026-02-01"})
    assert ww.json()["cny_per_usd"] is None


# ---- item codes on every export (all project types) ----

def test_wetworks_exports_qualify_names_and_codes():
    db = SessionLocal()
    try:
        line = _built_line(db, "wetworks", with_coverage=True, item_code="W1")
        line.product.odoo_id = "__export__.catalog_product"
        line.components[0].support_item.odoo_id = "__export__.catalog_si"
        db.commit()
        project, comp = line.project, line.components[0]
        assert comp.item_code is None  # blank component code falls back to the project code

        header, data = _sheet_rows(export_excel.build_sale_estimation_workbook(db, project))
        col = header.index
        c = "estimation_line_ids/sale_estimation_component_product_line_ids/"
        assert data[0][col("estimation_line_ids/product_id")] == f"{line.product.name} PJ W1"
        assert data[0][col("estimation_line_ids/default_code")] == "PJ W1"
        assert data[0][col(c + "product_id")] == f"{comp.support_item.name} PJ"
        assert data[0][col(c + "default_code")] == "PJ"
        # catalog odoo ids never ride along with project-specific names
        assert data[0][col("estimation_line_ids/product_id/id")] in (None, "")
        assert data[0][col(c + "product_id/id")] in (None, "")

        bh, bd = _sheet_rows(export_excel.build_bom_workbook(db, project))
        assert bd[0][bh.index("product")] == f"{line.product.name} PJ W1"
        assert bd[0][bh.index("reference")] == "PJ W1"
        assert bd[0][bh.index("bom_line_ids/product_id")] == f"{comp.support_item.name} PJ"

        sc = models.SellingCompany(name="SC wtest")
        db.add(sc)
        db.flush()
        project.selling_company_id = sc.id
        db.commit()
        _, rows = _sheet_rows(dict(export_excel.build_product_import_workbooks(db, project))["SC wtest"])
        by_name = {r[1]: r for r in rows}
        assert by_name[f"{line.product.name} PJ W1"][2] == "PJ W1"
        assert by_name[f"{comp.support_item.name} PJ"][2] == "PJ"
        assert all(r[0] in (None, "") for r in rows)
    finally:
        db.close()


def test_blank_item_code_falls_back_to_project_code():
    db = SessionLocal()
    try:
        line = _built_line(db, "wetworks", item_code=None)
        header, data = _sheet_rows(export_excel.build_sale_estimation_workbook(db, line.project))
        assert data[0][header.index("estimation_line_ids/default_code")] == "PJ"
        assert data[0][header.index("estimation_line_ids/product_id")] == f"{line.product.name} PJ"
    finally:
        db.close()


def test_export_blocked_without_project_code(client):
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    pid = _new_project(client, token, "Legacy no code")
    db = SessionLocal()
    try:
        db.query(models.Project).get(pid).code = None  # a project created before codes were required
        db.commit()
    finally:
        db.close()
    r = client.get(f"/api/projects/{pid}/export/bom", headers=h)
    assert r.status_code == 422 and "project code" in r.json()["detail"]


def _save(client, h, pid, locations, lines):
    return client.put(f"/api/projects/{pid}/estimate", headers=h,
                      json={"locations": locations, "lines": lines})


def test_furniture_same_product_and_code_must_share_bom(client):
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    prod, si, pid, loc = _furniture_project(client, h, _country_id(client, token))
    line = {"product_id": prod, "qty": 1, "item_code": "CH-1", "factory_work_cost_cny": 50}
    r = _save(client, h, pid, [{"id": loc, "name": "L"}, {"key": "k2", "name": "L2"}],
              [{**line, "location_id": loc}, {**line, "location_key": "k2"}])
    assert r.status_code == 200, r.text
    l1, l2 = (l["id"] for l in r.json()["lines"])
    for lid in (l1, l2):
        assert client.post(f"/api/estimate-lines/{lid}/components", headers=h, json={
            "support_item_id": si, "qty_per_unit": 2, "unit_price_cny": 30, "item_code": "FAB"}).status_code == 200

    # identical BOM -> exports, and the BOM sheet has a single mrp.bom for the pair
    r = client.get(f"/api/projects/{pid}/export/bom", headers=h)
    assert r.status_code == 200
    import io
    bh, bd = _sheet_rows(io.BytesIO(r.content))
    assert [row[bh.index("product")] for row in bd if row[bh.index("product")]] == ["F chair FP CH-1"]

    # change one line's quantity -> same product + code, different BOM -> blocked
    comp = next(l for l in client.get(f"/api/projects/{pid}/estimate-lines", headers=h).json()
                if l["id"] == l2)["components"][0]
    client.put(f"/api/estimate-line-components/{comp['id']}", headers=h, json={
        "support_item_id": si, "qty_per_unit": 3, "unit_price_cny": 30, "item_code": "FAB"})
    r = client.get(f"/api/projects/{pid}/export/sale-estimation", headers=h)
    assert r.status_code == 422 and "same BOM" in r.json()["detail"]


# ---- whole-estimate save + workspace ----

def _wetworks_catalog():
    db = SessionLocal()
    try:
        _line_seq[0] += 1
        ww = models.Product(name=f"WW save {_line_seq[0]}", uom="Sqm", category="Tile", product_type="wetworks")
        lf = models.Product(name=f"LF save {_line_seq[0]}", uom="Pcs", category="Seating",
                            product_type="loose_furniture")
        db.add_all([ww, lf])
        db.commit()
        return ww.id, lf.id
    finally:
        db.close()


def test_save_estimate_creates_updates_deletes(client):
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    ww, _ = _wetworks_catalog()
    pid = _new_project(client, token, "Save test")
    r = _save(client, h, pid, [{"key": "a", "name": "Lobby"}, {"key": "b", "name": "Rooms"}],
              [{"location_key": "a", "product_id": ww, "qty": 10},
               {"location_key": "b", "product_id": ww, "qty": 5, "item_code": "T-1"}])
    assert r.status_code == 200, r.text
    body = r.json()
    lobby = next(l for l in body["project"]["locations"] if l["name"] == "Lobby")
    keep, drop = body["lines"]

    # update one line, drop the other, drop "Rooms", add a location
    r = _save(client, h, pid, [{"id": lobby["id"], "name": "Main Lobby"}, {"key": "c", "name": "Spa"}],
              [{"id": keep["id"], "location_id": lobby["id"], "product_id": ww, "qty": 12},
               {"location_key": "c", "product_id": ww, "qty": 1}])
    assert r.status_code == 200, r.text
    ws = client.get(f"/api/projects/{pid}/workspace", headers=h).json()
    assert sorted(l["name"] for l in ws["project"]["locations"]) == ["Main Lobby", "Spa"]
    assert sorted(l["qty"] for l in ws["lines"]) == [1, 12]
    assert drop["id"] not in {l["id"] for l in ws["lines"]}
    assert ws["products"] and ws["support_items"] is not None and ws["selling_companies"] is not None
    lean = client.get(f"/api/projects/{pid}/workspace?catalog=false", headers=h).json()
    assert lean["products"] is None and len(lean["lines"]) == 2


def test_save_estimate_is_all_or_nothing(client):
    token = _login(client, "alice", "alicepass")
    h = _auth(token)
    ww, lf = _wetworks_catalog()
    pid = _new_project(client, token, "Atomic test")
    r = _save(client, h, pid, [{"key": "a", "name": "Lobby"}], [{"location_key": "a", "product_id": ww, "qty": 3}])
    assert r.status_code == 200
    before = client.get(f"/api/projects/{pid}/workspace?catalog=false", headers=h).json()
    loc_id = before["project"]["locations"][0]["id"]

    # a good new location + line, then a furniture product on a wetworks project -> rejected
    r = _save(client, h, pid, [{"id": loc_id, "name": "Lobby"}, {"key": "n", "name": "New"}],
              [{"id": before["lines"][0]["id"], "location_id": loc_id, "product_id": ww, "qty": 99},
               {"location_key": "n", "product_id": ww, "qty": 1},
               {"location_key": "n", "product_id": lf, "qty": 1}])
    assert r.status_code == 422
    after = client.get(f"/api/projects/{pid}/workspace?catalog=false", headers=h).json()
    assert [l["name"] for l in after["project"]["locations"]] == ["Lobby"]
    assert [(l["id"], l["qty"]) for l in after["lines"]] == [(before["lines"][0]["id"], 3)]


def test_save_estimate_rejects_other_projects_rows(client):
    alice, bob = _login(client, "alice", "alicepass"), _login(client, "bob", "bobpass")
    ww, _ = _wetworks_catalog()
    a_pid = _new_project(client, alice, "Alice rows")
    r = _save(client, _auth(alice), a_pid, [{"key": "a", "name": "L"}], [{"location_key": "a", "product_id": ww, "qty": 1}])
    a_loc = r.json()["project"]["locations"][0]["id"]
    b_pid = _new_project(client, bob, "Bob rows")
    r = _save(client, _auth(bob), b_pid, [{"id": a_loc, "name": "stolen"}], [])
    assert r.status_code == 422
    assert client.put(f"/api/projects/{a_pid}/estimate", headers=_auth(bob),
                      json={"locations": [], "lines": []}).status_code == 404


# ---- countries: new ones copy Saudi Arabia ----

def test_new_country_copies_ksa_rate_card_and_prices(client):
    db = SessionLocal()
    try:
        ksa = models.Country(name="Saudi Arabia", code="KSA", is_template=False,
                             material_currency_code="SAR", material_fx_rate_to_usd=3.75,
                             working_days_per_month=26, food_per_month_local=900, wages_oh_rate_pct=0.2)
        si = models.SupportItem(name="KSA cement", uom="Bag")
        db.add_all([ksa, si])
        db.flush()
        db.add(models.CountryMaterialPrice(country_id=ksa.id, support_item_id=si.id, unit_price_local=18.5))
        db.commit()
        si_id = si.id
    finally:
        db.close()

    token = _login(client, "alice", "alicepass")
    r = client.post("/api/countries", headers=_auth(token), json={"name": "Nigeria", "code": "NGA"})
    assert r.status_code == 200, r.text
    c = r.json()
    assert (c["name"], c["code"], c["is_template"]) == ("Nigeria", "NGA", True)
    assert c["material_currency_code"] == "SAR" and c["material_fx_rate_to_usd"] == 3.75
    assert c["food_per_month_local"] == 900 and c["wages_oh_rate_pct"] == 0.2
    prices = client.get(f"/api/countries/{c['id']}/material-prices", headers=_auth(token)).json()
    assert [(p["support_item_id"], p["unit_price_local"]) for p in prices] == [(si_id, 18.5)]

    # deleting: admin only, never the source country, never one in use
    assert client.delete(f"/api/countries/{c['id']}", headers=_auth(token)).status_code == 403
    admin = _auth(_login(client, "admin", "adminpass"))
    ksa_id = next(x["id"] for x in client.get("/api/countries", headers=admin).json() if x["code"] == "KSA")
    assert client.delete(f"/api/countries/{ksa_id}", headers=admin).status_code == 400
    assert client.delete(f"/api/countries/{_country_id(client, token)}", headers=admin).status_code == 400
    assert client.delete(f"/api/countries/{c['id']}", headers=admin).status_code == 200
    assert client.get(f"/api/countries/{c['id']}/material-prices", headers=admin).json() == []
